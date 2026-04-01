"""Excel (.xlsx) 文件脱敏处理器

处理策略：
- 遍历所有工作表的所有行
- 对相邻单元格进行键值对扫描
- 对每个单元格进行正则扫描
"""
from openpyxl import load_workbook

from processors.base import BaseProcessor
from core.scanner import Scanner
from core.logger import RedactionLog


class ExcelProcessor(BaseProcessor):
    """Excel 文件处理器"""

    def process(self, input_path: str, output_path: str, log: RedactionLog) -> bool:
        wb = load_workbook(input_path)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            self._process_sheet(ws, sheet_name, log)

        wb.save(output_path)
        return True

    def _process_sheet(self, ws, sheet_name: str, log: RedactionLog):
        """处理单个工作表"""
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            cells = list(row)

            # 1. 键值对扫描（相邻单元格作为 key-value）
            for i in range(0, len(cells) - 1):
                key_cell = cells[i]
                value_cell = cells[i + 1]

                key = str(key_cell.value or '').strip()
                value = str(value_cell.value or '').strip()

                if not key or not value:
                    continue

                kv_matches = self.scanner.scan_key_value_pairs([(key, value)])
                for m in kv_matches:
                    location = f'{sheet_name}!{value_cell.coordinate}'
                    self._add_log_entry(log, m.rule_id, m.rule_name,
                                        m.value_text, location, m.match_type)
                    # 替换单元格值
                    cell_str = str(value_cell.value)
                    value_cell.value = cell_str.replace(
                        m.value_text, self.scanner.replacement)

            # 2. 对每个单元格做独立正则扫描
            for cell in cells:
                cell_value = str(cell.value or '').strip()
                if not cell_value:
                    continue

                matches = self.scanner.scan_text(cell_value)
                for m in matches:
                    location = f'{sheet_name}!{cell.coordinate}'
                    self._add_log_entry(log, m.rule_id, m.rule_name,
                                        m.matched_text, location, m.match_type)
                    cell_str = str(cell.value)
                    cell.value = cell_str.replace(
                        m.matched_text, self.scanner.replacement)
