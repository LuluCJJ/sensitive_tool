"""Word (.docx) 文件脱敏处理器

处理策略：
- 遍历文档中的所有段落和表格
- 对段落文本进行正则 + 关键字扫描，直接替换文本
- 对表格按键值对逻辑扫描（标签→值）
"""
from docx import Document

from processors.base import BaseProcessor
from core.scanner import Scanner
from core.logger import RedactionLog


class WordProcessor(BaseProcessor):
    """Word 文件处理器"""

    def process(self, input_path: str, output_path: str, log: RedactionLog) -> bool:
        doc = Document(input_path)

        # 处理所有段落
        for i, para in enumerate(doc.paragraphs):
            if not para.text.strip():
                continue
            self._process_paragraph(para, f'段落{i+1}', log)

        # 处理所有表格
        for t_idx, table in enumerate(doc.tables):
            self._process_table(table, f'表格{t_idx+1}', log)

        doc.save(output_path)
        return True

    def _process_paragraph(self, para, location: str, log: RedactionLog):
        """处理单个段落"""
        text = para.text
        matches = self.scanner.scan_text(text)
        if not matches:
            return

        # 记录日志
        for m in matches:
            self._add_log_entry(log, m.rule_id, m.rule_name,
                                m.matched_text, location, m.match_type)

        # 替换段落中的文字（需要处理 runs 以保持格式）
        redacted_text = self.scanner.redact_text(text, matches)
        self._replace_paragraph_text(para, text, redacted_text)

    def _process_table(self, table, location_prefix: str, log: RedactionLog):
        """处理表格

        两种扫描策略：
        1. 键值对扫描：将相邻单元格作为 key-value 对
        2. 单元格内文本正则扫描
        """
        for r_idx, row in enumerate(table.rows):
            cells = row.cells
            # 键值对扫描（偶数列为key，奇数列为value）
            for c_idx in range(0, len(cells) - 1, 2):
                key = cells[c_idx].text.strip()
                value_cell = cells[c_idx + 1]
                value = value_cell.text.strip()

                if key and value:
                    kv_matches = self.scanner.scan_key_value_pairs([(key, value)])
                    for m in kv_matches:
                        location = f'{location_prefix} 第{r_idx+1}行第{c_idx+2}列'
                        self._add_log_entry(log, m.rule_id, m.rule_name,
                                            m.value_text, location, m.match_type)
                        # 替换单元格文本
                        self._replace_cell_text(value_cell, m.value_text,
                                                self.scanner.replacement)

            # 对每个单元格做独立的正则扫描
            for c_idx, cell in enumerate(cells):
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                matches = self.scanner.scan_text(cell_text)
                for m in matches:
                    location = f'{location_prefix} 第{r_idx+1}行第{c_idx+1}列'
                    self._add_log_entry(log, m.rule_id, m.rule_name,
                                        m.matched_text, location, m.match_type)
                    self._replace_cell_text(cell, m.matched_text,
                                            self.scanner.replacement)

    def _replace_paragraph_text(self, para, old_text: str, new_text: str):
        """替换段落文本，尽量保持原有格式"""
        if old_text == new_text:
            return

        # 简单策略：清空所有runs，在第一个run中写入新文本
        if para.runs:
            # 保持第一个run的格式
            first_run = para.runs[0]
            first_run.text = new_text
            for run in para.runs[1:]:
                run.text = ''
        else:
            para.text = new_text

    def _replace_cell_text(self, cell, old_value: str, new_value: str):
        """替换单元格中的特定文本"""
        for para in cell.paragraphs:
            if old_value in para.text:
                full_text = para.text.replace(old_value, new_value)
                if para.runs:
                    para.runs[0].text = full_text
                    for run in para.runs[1:]:
                        run.text = ''
                else:
                    para.text = full_text
